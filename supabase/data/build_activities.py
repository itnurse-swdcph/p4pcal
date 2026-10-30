"""Build the flat P4P activity catalog from the seven role worksheets.

Usage: python3 build_activities.py <source.xlsx> <output.json>
"""
import json
import re
import sys

import openpyxl


SHEETS = [
    ("พยาบาล", "พยาบาล", "NUR"),
    ("พนักงานช่วยเหลือคนไข้", "ผู้ช่วยเหลือคนไข้", "AID"),
    ("พนักงานประจำตึก", "พนักงานประจำตึก", "WARD"),
    ("พนักงานธุรการ", "พนักงานธุรการ", "ADMIN"),
    ("พนักงานเปล", "พนักงานเปล", "STRETCHER"),
    ("Paramedics", "Paramedic", "PMD"),
    ("EMTเจ้าพนักงานฉุกเฉินการแพทย์", "EMT", "EMT"),
]


def clean(value):
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def build(source):
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    activities = []
    for sheet_name, eligible_line, prefix in SHEETS:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Missing worksheet: {sheet_name}")
        sheet = workbook[sheet_name]
        index = 0
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, max_col=6, values_only=True), start=2):
            title = clean(row[1])
            uom = clean(row[2])
            minutes = row[3]
            if not title or not uom or minutes is None:
                continue
            try:
                minutes = float(minutes)
            except (TypeError, ValueError) as error:
                raise ValueError(f"Invalid minutes at {sheet_name}!D{row_number}: {minutes!r}") from error
            if minutes <= 0:
                raise ValueError(f"Minutes must be positive at {sheet_name}!D{row_number}")
            index += 1
            # The four medical-gas activities are appended to the aide sheet and
            # are visible only to staff assigned to the medical-equipment center.
            center_only = sheet_name == "พนักงานช่วยเหลือคนไข้" and row_number in range(63, 67)
            code = f"{prefix}-MED-{index:03d}" if center_only else f"{prefix}-{index:03d}"
            activities.append({
                "code": code,
                "domain": "บริการ",  # Required by the current database schema; reporting is flat.
                "title": title,
                "uom": uom,
                "minutes": minutes,
                "use_role_mult": False,
                "max_day": None,
                "max_month": None,
                "need_evidence": False,
                "definition": None,
                "no_dup_group": None,
                "eligible": [eligible_line],
                "note": None,
                "active": True,
            })
    codes = [item["code"] for item in activities]
    if len(codes) != len(set(codes)):
        raise ValueError("Generated duplicate activity codes")
    return activities


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: python3 build_activities.py <source.xlsx> <output.json>")
    result = build(sys.argv[1])
    with open(sys.argv[2], "w", encoding="utf-8") as output:
        json.dump(result, output, ensure_ascii=False, indent=2)
        output.write("\n")
    print(f"Built {len(result)} activities from {len(SHEETS)} worksheets")
